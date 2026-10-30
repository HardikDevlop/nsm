from __future__ import annotations

from io import BytesIO

from backend.services.report_csv import spreadsheet_safe


def generate_report_xlsx(rows: list[dict[str, object]]) -> bytes:
    """Serialize scheduled report rows as a readable XLSX workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    headers = list(rows[0].keys()) if rows else ["device_id", "hostname"]
    sheet.append(headers)
    for row in rows:
        values = []
        for header in headers:
            value = row.get(header)
            if header == "downtime_seconds":
                total = max(0, round(float(value or 0)))
                days, remainder = divmod(total, 86400)
                hours, remainder = divmod(remainder, 3600)
                minutes, seconds = divmod(remainder, 60)
                value = f"{days}d {hours}h {minutes}m {seconds}s"
            values.append(spreadsheet_safe(value))
        sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(32, max(12, len(str(header)) + 2))

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
