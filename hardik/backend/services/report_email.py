"""Scheduled delivery of the previous calendar day's management report."""
from __future__ import annotations

from email.message import EmailMessage
from io import BytesIO
import logging
import smtplib
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from apscheduler.triggers.cron import CronTrigger
from sqlalchemy.orm import Session

from backend.config.settings import get_settings
from backend.database.session import SessionLocal
from backend.schemas.nms import ReportManagementFilters, ReportManagementSummary

logger = logging.getLogger(__name__)
DAILY_REPORT_JOB_ID = "report-management:yesterday-email"


def _timezone() -> ZoneInfo:
    name = get_settings().report_timezone
    try:
        return ZoneInfo(name)
    except ZoneInfoNotFoundError:
        logger.error("Unknown report timezone %s; using Asia/Kolkata", name)
        return ZoneInfo("Asia/Kolkata")


def resolve_report_recipients() -> list[str]:
    configured = get_settings().report_email_recipients
    return list(dict.fromkeys(value.strip() for value in configured.split(",") if value.strip()))


def report_email_status(_: Session | None = None) -> dict[str, object]:
    settings = get_settings()
    return {
        "enabled": settings.report_email_enabled,
        "schedule": f"{settings.report_email_hour:02d}:{settings.report_email_minute:02d}",
        "timezone": settings.report_timezone,
        "report_period": "yesterday",
        "format": "xlsx",
        "recipient_count": len(resolve_report_recipients()),
        "smtp_configured": bool(settings.smtp_host and (settings.smtp_from or settings.smtp_username)),
    }


def _duration(seconds: int | float | None) -> str:
    total = max(0, round(float(seconds or 0)))
    days, remainder = divmod(total, 86400)
    hours, remainder = divmod(remainder, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{days}d {hours}h {minutes}m {secs}s"


def build_report_workbook(summary: ReportManagementSummary) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Summary"
    overview.append(["Yesterday Management Report", None])
    overview.append(["Period Start", summary.period_start.isoformat(sep=" ")])
    overview.append(["Period End", summary.period_end.isoformat(sep=" ")])
    overview.append(["Total Devices", summary.total_devices])
    overview.append(["Availability %", summary.availability_pct])
    overview.append(["Total Downtime", _duration(summary.downtime_seconds)])
    overview.append(["Average SNMP Health", summary.avg_snmp_health if summary.avg_snmp_health is not None else "N/A"])
    overview.append(["Average Performance", summary.avg_performance_score if summary.avg_performance_score is not None else "N/A"])
    overview.append(["SLA Met %", summary.sla_met_pct])
    overview["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="2563EB")
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 24

    sheet = workbook.create_sheet("Devices")
    headers = [
        "Device", "IP Address", "Site", "Protocol", "Availability %", "Downtime",
        "SNMP Health", "SNMP Success %", "Performance", "Interfaces", "Interfaces Down",
        "Avg CPU %", "Avg Memory %", "Avg Latency ms", "Packet Loss %", "SLA",
    ]
    sheet.append(headers)
    for record in summary.records:
        sheet.append([
            record.hostname, record.ip_address, record.site_name or "", record.protocol,
            record.availability_pct, _duration(record.downtime_seconds), record.snmp_health,
            record.snmp_success_rate, record.performance_score, record.interface_count,
            record.interface_down_count, record.avg_cpu_percent, record.avg_memory_percent,
            record.avg_latency_ms, record.packet_loss_pct, record.sla_status,
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(28, max(12, len(header) + 2))

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def send_yesterday_report_email() -> bool:
    settings = get_settings()
    if not settings.report_email_enabled:
        logger.info("Daily report email is disabled")
        return False
    recipients = resolve_report_recipients()
    if not recipients:
        logger.warning("Daily report email skipped: REPORT_EMAIL_RECIPIENTS is empty")
        return False
    if not settings.smtp_host or not (settings.smtp_from or settings.smtp_username):
        logger.warning("Daily report email skipped: SMTP is not fully configured")
        return False

    with SessionLocal() as db:
        from backend.api.routes import _build_report_rows
        summary, _ = _build_report_rows(db, ReportManagementFilters(period="yesterday", protocol="all"))

    attachment = build_report_workbook(summary)
    report_date = summary.period_start.date().isoformat()
    message = EmailMessage()
    message["Subject"] = f"NMS Yesterday Report - {report_date}"
    message["From"] = settings.smtp_from or settings.smtp_username
    message["To"] = ", ".join(recipients)
    message.set_content(
        f"Attached is the NMS management report for {report_date}.\n\n"
        f"Devices: {summary.total_devices}\n"
        f"Availability: {summary.availability_pct:.2f}%\n"
        f"SLA met: {summary.sla_met_pct:.2f}%\n"
    )
    message.add_attachment(
        attachment,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"NMS_Yesterday_Report_{report_date}.xlsx",
    )
    try:
        with smtplib.SMTP(settings.smtp_host, settings.smtp_port, timeout=30) as smtp:
            if settings.smtp_use_tls:
                smtp.starttls()
            if settings.smtp_username:
                smtp.login(settings.smtp_username, settings.smtp_password)
            smtp.send_message(message)
        logger.info("Yesterday report emailed to %d configured recipient(s) for %s", len(recipients), report_date)
        return True
    except Exception:
        logger.exception("Could not email yesterday management report for %s", report_date)
        return False


def register_daily_report_job(polling_scheduler: object | None) -> bool:
    scheduler = getattr(polling_scheduler, "scheduler", None)
    if scheduler is None or not hasattr(scheduler, "add_job"):
        return False
    settings = get_settings()
    if not settings.report_email_enabled:
        existing = scheduler.get_job(DAILY_REPORT_JOB_ID) if hasattr(scheduler, "get_job") else None
        if existing and hasattr(scheduler, "remove_job"):
            scheduler.remove_job(DAILY_REPORT_JOB_ID)
        return False
    scheduler.add_job(
        send_yesterday_report_email,
        CronTrigger(hour=settings.report_email_hour, minute=settings.report_email_minute, timezone=_timezone()),
        id=DAILY_REPORT_JOB_ID,
        replace_existing=True,
        max_instances=1,
        coalesce=True,
        misfire_grace_time=3600,
    )
    logger.info("Yesterday-report email scheduled for %02d:%02d %s", settings.report_email_hour, settings.report_email_minute, settings.report_timezone)
    return True
